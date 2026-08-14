from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "manual_assets"
OUTPUT = ROOT / "output" / "manual" / "Tarangini_Workflow_Suite_1.1.6_User_Manual.docx"

NAVY = "17365D"
BLUE = "2E74B5"
LIGHT_BLUE = "EAF2F8"
PALE_GREEN = "EAF4EA"
PALE_YELLOW = "FFF4CE"
PALE_RED = "FDE9E7"
LIGHT_GRAY = "F2F4F7"
DARK = "243447"
MID = "5B6770"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BLUE, size="10", edge="left"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    tag = qn(f"w:{edge}")
    border = borders.find(tag)
    if border is None:
        border = OxmlElement(f"w:{edge}")
        borders.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MID)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color, before, after in (
        ("Title", 30, NAVY, 0, 14),
        ("Subtitle", 14, MID, 0, 10),
        ("Heading 1", 18, NAVY, 4, 10),
        ("Heading 2", 13.5, BLUE, 12, 6),
        ("Heading 3", 11.5, NAVY, 9, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(10.5)
    bullet.paragraph_format.left_indent = Inches(0.23)
    bullet.paragraph_format.first_line_indent = Inches(-0.16)
    bullet.paragraph_format.space_after = Pt(3)

    numbered = doc.styles["List Number"]
    numbered.font.name = "Calibri"
    numbered.font.size = Pt(10.5)
    numbered.paragraph_format.left_indent = Inches(0.28)
    numbered.paragraph_format.first_line_indent = Inches(-0.2)
    numbered.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = "TARANGINI WORKFLOW SUITE 1.1.6  |  VISUAL USER MANUAL"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.runs[0]
    header_run.font.name = "Calibri"
    header_run.font.size = Pt(8)
    header_run.font.bold = True
    header_run.font.color.rgb = RGBColor.from_string(MID)

    add_page_number(section.footer.paragraphs[0])


def add_callout(doc, title, text, fill=LIGHT_BLUE, border=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_width(cell, 9680)
    set_cell_shading(cell, fill)
    set_cell_border(cell, border, "12", "left")
    set_cell_margins(cell, 105, 150, 105, 150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title.upper())
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(border)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.08
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_steps(doc, steps):
    numbering = doc.part.numbering_part.element
    style_num_id = doc.styles["List Number"]._element.pPr.numPr.numId.val
    style_num = next(
        node for node in numbering.findall(qn("w:num"))
        if int(node.get(qn("w:numId"))) == int(style_num_id)
    )
    abstract_num_id = style_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    existing_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    new_num_id = max(existing_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_num_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    for step in steps:
        p = doc.add_paragraph(style="List Number")
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = new_num_id
        p.add_run(step)


def add_bullets(doc, bullets):
    for bullet in bullets:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(bullet)


def add_picture(doc, filename, caption, width=6.78):
    path = ASSETS / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(7)
    cap.paragraph_format.keep_with_next = True
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(MID)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        set_cell_width(cell, int(widths[index] * 1440))
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            set_cell_width(cell, int(widths[index] * 1440))
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if row_index % 2:
                set_cell_shading(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            r.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def new_page(doc, title, intro=None):
    doc.add_page_break()
    doc.add_heading(title, level=1)
    if intro:
        p = doc.add_paragraph(intro)
        p.paragraph_format.space_after = Pt(8)


def add_cover(doc):
    for _ in range(2):
        doc.add_paragraph()
    logo = ROOT / "assets" / "tarangini.png"
    if logo.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(logo), width=Inches(1.25))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("TARANGINI BILLING")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Visual User Manual")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Billing, payments, POS, corrections, backups and recovery")

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    line.paragraph_format.space_before = Pt(10)
    r = line.add_run("VERSION 1.1.6  |  OWNER & OPERATOR GUIDE")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(NAVY)

    doc.add_paragraph()
    add_callout(
        doc,
        "Purpose",
        "Use this guide for daily billing, payment collection, invoice corrections, operator handover, backup setup, restoration and software updates.",
        fill=LIGHT_BLUE,
        border=BLUE,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    r = p.add_run("Prepared June 2026")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(MID)


def build():
    doc = Document()
    style_document(doc)
    add_cover(doc)

    new_page(doc, "Quick Start")
    doc.add_heading("Who uses what", level=2)
    add_table(
        doc,
        ["Role", "Daily responsibilities", "Restricted or approval actions"],
        [
            ("Owner / Admin", "Full billing, accounting, reports, backups, updates and acceptance.", "Approves corrections; restores backups; installs updates."),
            ("Operator", "Creates bills, receives permitted payments, requests corrections and submits shift handover.", "Cannot approve own correction requests or perform protected administration."),
            ("Client counter", "Connects to the Main System over LAN and can keep offline drafts during a temporary outage.", "Must sync pending work before an upgrade."),
        ],
        [1.25, 3.15, 2.35],
    )
    doc.add_heading("Recommended opening routine", level=2)
    add_steps(
        doc,
        [
            "Start the Main System first. It holds the shared database and serves all client counters.",
            "Confirm every client displays Connected before billing.",
            "Open Dashboard and review today’s sales, payment status and pending work.",
            "At day end, submit operator collections and confirm the automatic backup completed.",
        ],
    )
    add_callout(
        doc,
        "Payment rule",
        "A Bank, UPI or Card selection on an invoice does not post money directly into the bank ledger. The bank ledger changes only when a linked Receipt Voucher is created.",
        fill=PALE_YELLOW,
        border="B7791F",
    )

    new_page(doc, "1. Main System and Client Setup")
    add_picture(doc, "12-main-client-setup.png", "Main System / Client selection and LAN connection setup.")
    add_steps(
        doc,
        [
            "On the server computer, choose Main System. Allow Windows Firewall access when prompted.",
            "On each counter, choose Client System and enter the Main System LAN IP and port shown by the server.",
            "Use Test Connection, then save. A connected client uses the Main System data.",
            "Keep the Main System computer on while counters are billing.",
        ],
    )
    add_callout(
        doc,
        "Backup ownership",
        "Configure automatic backups on the Main System. Client counters store only local connection details and offline drafts; they are not the authoritative database.",
        fill=PALE_GREEN,
        border="3B7A3E",
    )

    new_page(doc, "2. Dashboard and Navigation")
    add_picture(doc, "01-dashboard.png", "Dashboard summary with navigation available from the menu.")
    add_bullets(
        doc,
        [
            "Use Dashboard to review sales, collection mix and items needing attention.",
            "Owner screens include accounting, corrections, shift acceptance, backup and update management.",
            "Operator screens focus on billing, permitted receipts and handover reporting.",
            "Use the financial year selector before entering or reviewing older transactions.",
        ],
    )

    new_page(doc, "3. Create a Standard Sale Invoice")
    add_picture(doc, "02-sale-bill.png", "Sale invoice entry with payment, delivery and per-row description fields.")
    doc.add_heading("Rates include GST and round off", level=2)
    add_bullets(
        doc,
        [
            "For a Regular taxpayer, select Incl. GST when the entered item rate already contains GST. Tarangini back-calculates taxable value and tax.",
            "The inclusive-GST control is available on Sale, POS, Project Printing, Purchase, Expense, Credit/Debit Note, Return, Purchase Order and Job Estimate transactions.",
            "Round off final total to nearest rupee is available for every taxpayer type and is enabled by default.",
            "The signed rounding adjustment is stored, printed on the invoice when non-zero and posted to the Rounding Adjustment ledger.",
            "Clear the round-off checkbox only when the exact paise total must be retained.",
        ],
    )
    add_callout(
        doc,
        "Example",
        "An exclusive rate of Rs. 99.00 with 18% GST totals Rs. 116.82. With round off enabled, the invoice total is Rs. 117.00 and Round Off is Rs. 0.18.",
        fill=PALE_GREEN,
        border="3B7A3E",
    )
    add_steps(
        doc,
        [
            "Open Sale Bill and confirm company, voucher series, invoice date and customer.",
            "Select Cash, Credit, Bank/UPI/Card or Split as the payment mode.",
            "Add items by search, barcode scanner or manual barcode entry.",
            "Enter the invoice-row Description / Serial Number. This is separate from Item Master and is saved, printed, edited and duplicated with the invoice.",
            "Complete optional Delivery To details: address, recipient, phone, transporter/courier, tracking or LR number and dispatch date.",
            "Review tax and totals, then save the invoice.",
        ],
    )
    add_callout(
        doc,
        "Printing",
        "The printed invoice shows Bill To and Delivery To separately, includes row descriptions/serial numbers and lists the payment split.",
        fill=LIGHT_BLUE,
        border=BLUE,
    )

    new_page(doc, "4. Payments and Linked Receipts")
    add_picture(doc, "03-payment-received.png", "Payment received prompt shown after saving a non-cash invoice.")
    doc.add_heading("Bank, UPI and Card", level=2)
    add_steps(
        doc,
        [
            "Save the invoice with Bank/UPI/Card or a non-cash split portion.",
            "At Payment received?, choose Yes only when the payment is actually confirmed.",
            "Enter the transaction reference. Tarangini creates a linked Receipt Voucher.",
            "Choose No when payment is pending. Match and create the receipt later from the bank statement or receipt workflow.",
        ],
    )
    doc.add_heading("Split payment", level=2)
    add_bullets(
        doc,
        [
            "Cash portion posts to cash immediately.",
            "Bank, UPI and Card portions remain receivable until linked Receipt Vouchers are created.",
            "Invoice status becomes Paid, Partly Paid, Unpaid or Overdue according to settled value and due date.",
            "An overpayment remains visible for review; never hide it by changing the invoice total.",
        ],
    )

    new_page(doc, "5. POS Counter")
    add_picture(doc, "04-pos-counter.png", "Responsive POS screen for fast counter billing.")
    add_steps(
        doc,
        [
            "Search or scan an item, then adjust quantity and the row description/serial number if needed.",
            "Select an existing customer or use the permitted cash customer.",
            "Choose Cash, Bank/UPI/Card or Split and enter each component accurately.",
            "Save, answer the payment-received prompt for non-cash money, and print the invoice.",
        ],
    )
    add_callout(
        doc,
        "Barcode use",
        "USB scanners normally act like a keyboard. On supported phones/tablets, the LAN sales portal can use the camera; manual barcode entry remains available.",
        fill=LIGHT_BLUE,
        border=BLUE,
    )

    new_page(doc, "6. Invoice Payment Status")
    add_picture(doc, "05-invoice-status.png", "Invoice status list showing settlement and outstanding amounts.")
    add_table(
        doc,
        ["Status", "Meaning", "Action"],
        [
            ("Paid", "Linked receipts and cash equal or exceed the invoice total.", "No collection action."),
            ("Partly Paid", "Some value is settled; an amount remains receivable.", "Collect balance and create receipt."),
            ("Unpaid", "No qualifying settlement has been posted.", "Follow up or keep as credit."),
            ("Overdue", "Outstanding value remains after the due date.", "Prioritize collection."),
        ],
        [1.2, 3.1, 2.45],
    )
    add_callout(
        doc,
        "Reconciliation",
        "Use transaction references to match Receipt Vouchers with the bank statement. The invoice payment mode alone is not evidence that the bank received money.",
        fill=PALE_YELLOW,
        border="B7791F",
    )

    new_page(doc, "7. Invoice Correction Requests")
    add_picture(doc, "06-invoice-corrections.png", "Correction request queue for operator requests and owner decisions.")
    doc.add_heading("Operator", level=2)
    add_steps(
        doc,
        [
            "Open the invoice and request removal of the incorrect item or quantity.",
            "Enter a clear reason and submit. The original invoice remains in the audit history.",
        ],
    )
    doc.add_heading("Owner", level=2)
    add_steps(
        doc,
        [
            "Open Invoice Corrections and review invoice, item, quantity, tax and reason.",
            "Approve or reject with a note. Approval creates a linked Credit Note.",
            "Confirm stock, GST, customer balance and accounting reports reflect the Credit Note.",
        ],
    )
    add_callout(
        doc,
        "Do not delete",
        "Never erase an issued invoice to correct it. Use the correction request and linked Credit Note so the audit trail remains complete.",
        fill=PALE_RED,
        border="A61B1B",
    )

    new_page(doc, "8. Operator Collection and Handover")
    add_picture(doc, "07-shifts.png", "Shift report with collection breakdown and owner acceptance.")
    add_steps(
        doc,
        [
            "Select operator, counter, shift and date.",
            "Review Cash, Bank, UPI, Card, credit/pending and refunds separately.",
            "Enter actual cash handed over. The system calculates expected cash and variance.",
            "Include bank references and an operator note, then submit.",
            "Owner verifies the money and references, then accepts with date, time and note.",
        ],
    )
    add_callout(
        doc,
        "Variance",
        "A non-zero variance requires explanation before owner acceptance. Do not change invoice payments merely to force the handover to match.",
        fill=PALE_YELLOW,
        border="B7791F",
    )

    new_page(doc, "9. Backup Strategy")
    doc.add_heading("Use three layers", level=2)
    add_table(
        doc,
        ["Layer", "Purpose", "Recommended frequency"],
        [
            ("Automatic database copy", "Fast recovery from computer or database failure.", "Daily; keep in a synced or network folder."),
            ("Manual JSON export", "Portable business-data export for migration and audit.", "Weekly and before major data work."),
            ("Encrypted .tbe backup", "Complete disaster-recovery copy protected by a password.", "Monthly and before upgrades or migration."),
        ],
        [1.75, 3.25, 1.75],
    )
    doc.add_heading("Good backup practice", level=2)
    add_bullets(
        doc,
        [
            "Keep at least one copy outside the Main System computer.",
            "Use a OneDrive/Google Drive synced folder or a reliable network share for automatic backups.",
            "Keep the encrypted-backup password separately. It cannot be recovered from the backup file.",
            "Test restoration periodically on a spare system or test profile.",
            "Do not upgrade while client counters have pending offline invoices.",
        ],
    )
    add_callout(
        doc,
        "Important",
        "A backup stored only on the same disk as the live database does not protect against disk loss, theft or ransomware.",
        fill=PALE_RED,
        border="A61B1B",
    )

    new_page(doc, "10. Manual Backup and Restore")
    add_picture(doc, "08-backup-restore.png", "Backup & Restore tools for JSON export, import and encrypted recovery.")
    doc.add_heading("Create a JSON backup", level=2)
    add_steps(
        doc,
        [
            "Open Settings > Backup & Restore.",
            "Choose the organization and financial year, or export all data when preparing a full backup.",
            "Select Export Backup and store the downloaded JSON file in the approved backup location.",
        ],
    )
    doc.add_heading("Restore or migrate", level=2)
    add_bullets(
        doc,
        [
            "JSON Import validates the file and merges supported records into the current database.",
            "Encrypted Disaster-Recovery Restore replaces the live database and requires an application restart.",
            "Take a fresh backup before either operation and perform protected restores as Owner/Admin.",
        ],
    )

    new_page(doc, "11. Automatic Backup Settings")
    add_picture(doc, "09-automatic-backup-settings.png", "Automatic backup settings. The displayed folder is a documentation test example.")
    add_steps(
        doc,
        [
            "Open Settings > Automatic Backup.",
            "Enable automatic backup and choose an interval from 1 to 720 hours. Daily is recommended.",
            "Choose a local, network or cloud-synced folder writable by the Main System Windows account.",
            "Save Settings, then choose Run Backup Now to verify the folder immediately.",
            "Open the folder and confirm a tarangini-auto-<date-time>.db file was created.",
        ],
    )
    add_callout(
        doc,
        "Folder shown above",
        "The screenshot was captured in a documentation test profile. Your installed application uses the location configured on your own Main System.",
        fill=PALE_YELLOW,
        border="B7791F",
    )

    new_page(doc, "12. Backup and Data Locations")
    add_table(
        doc,
        ["Data", "Default Windows location", "Notes"],
        [
            ("Live database", r"%APPDATA%\Tarangini Billing\database\tarangini.db", "Authoritative Main System database."),
            ("Automatic backups", r"%APPDATA%\Tarangini Billing\database\backups", "Used when no custom folder is selected."),
            ("Pre-upgrade recovery", r"%APPDATA%\Tarangini Billing\database\upgrade-backups", "Created by Update Manager before installation."),
            ("Network settings", r"%APPDATA%\Tarangini Billing\network.json", "Main/Client role, host and port."),
            ("Client offline drafts", r"%APPDATA%\Tarangini Billing\offline-client", "Pending local drafts and sync queue."),
            ("Manual JSON/.tbe", "Your chosen folder; JSON normally downloads through the browser.", "Move to the approved backup location after export."),
        ],
        [1.55, 3.45, 1.75],
    )
    doc.add_heading("Open the location", level=2)
    add_steps(
        doc,
        [
            "Press Windows + R.",
            r"Enter %APPDATA%\Tarangini Billing and press Enter.",
            "Open database for the live DB, backups and upgrade-backups folders.",
        ],
    )
    add_callout(
        doc,
        "Recommended custom location",
        r"Example: D:\Tarangini Backups or a synced folder such as C:\Users\<user>\OneDrive\Tarangini Backups. Confirm the folder actually synchronizes and has sufficient free space.",
        fill=PALE_GREEN,
        border="3B7A3E",
    )

    new_page(doc, "13. Update Manager")
    add_picture(doc, "11-update-manager.png", "Update Manager showing installed and available package information.")
    add_steps(
        doc,
        [
            "Confirm all client counters are connected and pending offline invoices are zero.",
            "Create and verify a fresh backup.",
            "Open Update Manager and compare Current Version with Latest Version.",
            "Download from the Main System LAN distribution source. Tarangini verifies the SHA-256 checksum.",
            "Install only after verification succeeds. The updater preserves the database, client code, server IP and pending sync state.",
            "After restart, verify the version, open Dashboard and test one read-only report before normal billing.",
        ],
    )
    add_callout(
        doc,
        "Unsafe upgrade block",
        "The update must be postponed when offline invoices are pending. Sync them first so the server and clients do not diverge.",
        fill=PALE_RED,
        border="A61B1B",
    )

    doc.add_heading("HTTPS-ready web deployment", level=2)
    add_bullets(
        doc,
        [
            "For internet or cloud testing, place Tarangini behind a trusted HTTPS reverse proxy, or configure its PEM/PFX certificate environment variables.",
            "Camera barcode access and PWA installation on phones require HTTPS, except on localhost.",
            "Keep the database and automatic backup folder on persistent storage and restrict CORS to approved origins.",
            "Do not expose the application directly to the public internet with recovery passwords unchanged.",
        ],
    )

    new_page(doc, "14. Mobile Sales and Offline Sync")
    doc.add_heading("LAN mobile portal", level=2)
    add_bullets(
        doc,
        [
            "Connect the phone/tablet to the same trusted LAN or Wi-Fi as the Main System.",
            "Open the Main System LAN address supplied by the administrator.",
            "Sign in with the assigned role. Search/browse items, select customer and create Cash, Bank or Split bills.",
            "Use camera barcode scanning when the browser permits it; otherwise use manual entry or a paired scanner.",
            "The same payment and audit rules apply as on desktop.",
        ],
    )
    doc.add_heading("Offline drafts", level=2)
    add_bullets(
        doc,
        [
            "Offline mode is for temporary connection loss, not routine operation.",
            "A pending draft must remain visible until the Main System confirms sync.",
            "Do not uninstall, clear application data, change server identity or install an update while drafts are pending.",
            "After reconnection, run sync and verify the invoice appears on the Main System.",
        ],
    )
    add_callout(
        doc,
        "Client mismatch",
        "When a client-version mismatch warning appears, stop new billing on that client, finish or sync pending drafts, then update it from the Main System package.",
        fill=PALE_YELLOW,
        border="B7791F",
    )

    new_page(doc, "15. Daily Controls and Troubleshooting")
    doc.add_heading("Day-end checklist", level=2)
    add_bullets(
        doc,
        [
            "All invoices show the expected Paid, Partly Paid, Unpaid or Overdue status.",
            "Bank/UPI/Card receipts have transaction references and match the statement or terminal report.",
            "Operator handover is submitted; owner acceptance and variance note are recorded.",
            "Client offline queues are empty.",
            "Today’s automatic backup file exists outside the live database folder when possible.",
        ],
    )
    doc.add_heading("Common problems", level=2)
    add_table(
        doc,
        ["Problem", "Check first", "Resolution"],
        [
            ("Client cannot connect", "Main System running, same LAN, correct IP/port, firewall.", "Use Test Connection and correct network settings."),
            ("Bank balance did not change", "Is there a linked Receipt Voucher?", "Create/match the receipt with transaction reference."),
            ("Automatic backup missing", "Enabled, interval, path exists, write permission.", "Run Backup Now and correct the destination."),
            ("Restore rejected", "Correct file type, password and organization/year.", "Use a valid backup and Owner/Admin access."),
            ("Update blocked", "Pending offline drafts or failed checksum.", "Sync first; re-download a verified package."),
        ],
        [1.5, 2.6, 2.65],
    )
    add_callout(
        doc,
        "Before support",
        "Record the exact error, current version, Main/Client role, organization, financial year and the time the issue occurred. Do not send passwords or unencrypted customer data.",
        fill=LIGHT_BLUE,
        border=BLUE,
    )

    new_page(doc, "Quick Reference")
    add_table(
        doc,
        ["Task", "Menu / screen", "Key rule"],
        [
            ("Create invoice", "Sale Bill or POS", "Enter row description/serial and correct payment split."),
            ("Record bank money", "Payment prompt / Receipt Voucher", "Bank posts only through linked receipt."),
            ("Correct invoice", "Invoice Corrections", "Owner approval creates Credit Note."),
            ("Close shift", "Operator Collections / Shifts", "Explain cash variance; owner accepts."),
            ("Backup now", "Settings > Automatic Backup", "Run Backup Now and verify the file."),
            ("Export data", "Settings > Backup & Restore", "Store JSON/.tbe outside the Main System disk."),
            ("Upgrade", "Update Manager", "Backup, sync all drafts, verify SHA-256."),
        ],
        [1.5, 2.65, 2.6],
    )
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("END OF MANUAL")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(NAVY)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
