from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


workspace = Path(r"C:\Users\user\Documents\New project")
project = workspace / "tarangini billing - brave"
manual = workspace / "Tarangini Billing 2.3.0 Visual User Manual.docx"
screens = project / "output" / "responsive-chrome"

document = Document(manual)
document.add_page_break()
document.add_heading("15. PIN Login, Stock Journal and Mobile Access", level=1)
document.add_paragraph(
    "Version 2.3.0 uses a simple PIN for daily owner and operator access and adds "
    "built-in reference tools for stock and mobile billing."
)

document.add_heading("Owner1 assigns user PINs", level=2)
for text in (
    "Open Settings > Users as Owner1.",
    "Create or edit Owner2 and operator accounts, assign company access and permissions, "
    "then enter a unique 4-6 digit PIN.",
    "Use Reset PIN when a user forgets the PIN. Recovery passwords are retained only for "
    "migration or account recovery.",
):
    document.add_paragraph(text, style="List Bullet")

document.add_heading("Dashboard mobile operator QR", level=2)
document.add_paragraph(
    "Connect the phone or tablet to the same Wi-Fi/LAN. Scan the dashboard QR code, "
    "then sign in with the operator username and PIN. The QR stores only the server URL."
)
paragraph = document.add_paragraph()
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph.add_run().add_picture(str(screens / "desktop-dashboard.png"), width=Inches(6.25))

document.add_heading("Stock Movement Journal", level=2)
document.add_paragraph(
    "Open Inventory & Stock and scroll below the item summary. The Stock Movement Journal "
    "shows the date, source reference, transaction type, item, quantity in/out, rate and "
    "movement value for purchases, sales, returns and approved corrections."
)
paragraph = document.add_paragraph()
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph.add_run().add_picture(str(screens / "desktop-stock.png"), width=Inches(6.25))

document.add_heading("Preloaded Help & FAQ", level=2)
document.add_paragraph(
    "Open Help & FAQ from the System menu. Search for PIN, payment, stock, GST, backup, "
    "offline sync, update manager or mobile access. The guide is included in the application "
    "and works without internet."
)
paragraph = document.add_paragraph()
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph.add_run().add_picture(str(screens / "desktop-help.png"), width=Inches(6.25))

document.save(manual)
